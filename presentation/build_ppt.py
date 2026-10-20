"""Bangun bukti_bigdata.pptx. Snippet kode dibaca langsung dari file sumber."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(__file__).resolve().parent / "bukti_bigdata.pptx"

BG, PANEL = RGBColor(0x0F, 0x17, 0x2A), RGBColor(0x1E, 0x29, 0x3B)
ACC, TXT, MUTED, CODE = (RGBColor(0x38, 0xBD, 0xF8), RGBColor(0xF1, 0xF5, 0xF9),
                         RGBColor(0x94, 0xA3, 0xB8), RGBColor(0xA7, 0xF3, 0xD0))

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


def snip(rel, start, end):
    lines = (ROOT / rel).read_text(encoding="utf-8").splitlines()[start - 1:end]
    return "\n".join(lines), f"{rel.replace('/', chr(92))}  (baris {start}-{end})"


def text(slide, x, y, w, h, s, size, color=TXT, bold=False, font="Calibri"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
        r.font.color.rgb = color
    return tb


def code_box(slide, x, y, w, h, body, caption, size):
    box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = PANEL; box.line.fill.background()
    text(slide, x + .1, y + .05, w - .2, .3, caption, 10, ACC, True)
    text(slide, x + .1, y + .35, w - .2, h - .4, body, size, CODE, font="Consolas")


def new_slide(title, sub=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    text(s, .5, .3, 12.3, .7, title, 28, TXT, True)
    if sub:
        text(s, .5, .95, 12.3, .4, sub, 14, MUTED)
    return s


def code_slide(title, sub, snippets, size=10, note=""):
    s = new_slide(title, sub)
    n = len(snippets)
    gap = .2
    w = (12.3 - gap * (n - 1)) / n
    h = 5.3 if not note else 4.9
    for i, (body, cap) in enumerate(snippets):
        code_box(s, .5 + i * (w + gap), 1.5, w, h, body, cap, size)
    if note:
        text(s, .5, 6.55, 12.3, .8, note, 13, MUTED)
    return s


def bullet_slide(title, sub, bullets, size=18):
    s = new_slide(title, sub)
    text(s, .7, 1.6, 11.9, 5.2, "\n".join("•  " + b for b in bullets), size)
    return s


# 1 Judul
s = new_slide("Bukti Tugas Pengelolaan Big Data", "Analisis Sentimen Publik: Putusan MK Isu Ijazah Gibran")
text(s, .7, 1.9, 11.9, 5, "\n".join([
    "Arsitektur Medallion di MongoDB:",
    "",
    "  YouTube API v3  ─┐",
    "                              ├─►  BRONZE (raw)  ─►  SILVER (bersih)  ─►  GOLD (sentimen)  ─►  Dashboard Streamlit",
    "  Scraping Antara ─┘",
    "",
    "Database: bigdata_sentiment_mk  |  Python 3.11  |  Docker Compose",
    "Slide ini hanya memuat bukti dari kode yang ada di repositori.",
]), 18, TXT, font="Consolas")

# 2 Container NoSQL
code_slide("1. Pembuatan Container NoSQL (MongoDB)", "docker-compose.yml: service mongodb + mongo-express",
           [snip("docker-compose.yml", 1, 28)], 9,
           "Image mongo:7.0, container pbd_mongodb, port 27017, data di volume mongo_data. Mongo Express (port 8081) untuk inspeksi data.")

# 3 YouTube
code_slide("2a. Sumber Data 1: YouTube Data API v3", "Semi-structured: JSON komentar dengan objek nested 'raw'",
           [snip("collectors/youtube_api.py", 41, 48), snip("collectors/youtube_api.py", 78, 104)], 9,
           "search().list mencari video; commentThreads().list mengambil komentar + balasan (paginasi nextPageToken).")

# 4 Scraping
code_slide("2b. Sumber Data 2: Web Scraping Berita (Antara News)", "Unstructured: HTML artikel → dokumen JSON, memakai Scrapling",
           [snip("collectors/news_scraper.py", 18, 36), snip("collectors/news_scraper.py", 49, 78)], 8,
           "Selector CSS untuk daftar artikel, lalu fetch halaman detail (judul, paragraf, penulis, tanggal).")

# 5 Raw
code_slide("3. Penyimpanan Raw Data (Bronze Layer)", "Data mentah disimpan tanpa modifikasi, deduplikasi lewat unique index",
           [snip("pipeline/bronze.py", 17, 27), snip("pipeline/bronze.py", 44, 57)], 10,
           "Koleksi: bronze_youtube (unique comment_id) dan bronze_news (unique url). $setOnInsert menjaga data mentah tetap utuh.")

# 6 Structured
code_slide("4. Pengolahan Raw → Data Terstruktur (Silver & Gold)", "Cleaning NLP → skema tetap → klasifikasi sentimen",
           [snip("pipeline/preprocessing.py", 124, 136), snip("pipeline/silver.py", 52, 72), snip("pipeline/gold.py", 61, 75)], 8,
           "Raw: {raw:{text,author,...}}  →  Silver: text_original/text_clean/author/like_count  →  Gold: + sentiment, confidence, probabilities.")

# 7 Angka
bullet_slide("Hasil Pengumpulan & Pengolahan", "Sumber angka: laporan_hasil_analisis_sentimen.md (bukan hasil run ulang)", [
    "YouTube: 2 video, 234 komentar  |  Berita Antara: 3 artikel",
    "bronze_youtube: 240  |  bronze_news: 3",
    "silver_comments: 234  |  silver_news: 3",
    "gold_sentiment: 234",
    "Sentimen: Positif 84 (35,9%)  |  Netral 76 (32,5%)  |  Negatif 74 (31,6%)",
    "Catatan: collector fallback ke data mock bila API/scraping gagal (collectors/mock_data.py).",
], 20)

# 8 Volume
code_slide("Bonus 1. Docker Volume untuk Penyimpanan Data", "Data MongoDB di named volume, terpisah dari filesystem container",
           [snip("docker-compose.yml", 10, 11), snip("docker-compose.yml", 50, 53)], 14,
           "Volume pbd_mongo_data (driver local): data bertahan saat container dihapus/di-restart. Catatan: named volume, bukan bind-mount ke disk lain.")

# 9 Scheduler
code_slide("Bonus 2. Crawling Otomatis (Scheduler)", "Siklus: collect → Bronze → Silver → Gold, berulang tiap N menit",
           [snip("scheduler/scheduler.py", 36, 55), snip("scheduler/scheduler.py", 64, 78)], 9,
           "Dijalankan: python scheduler/scheduler.py [menit]  (default 60 menit). Berjalan dari host, belum jadi service compose.")

# 10 Index
code_slide("Bonus 3. Indexing & Pencarian", "Index MongoDB per layer + pencarian kata kunci di dashboard",
           [snip("pipeline/gold.py", 21, 28), snip("pipeline/silver.py", 21, 27), snip("dashboard/app.py", 381, 390)], 9,
           "Unique index (comment_id, url) + index sentiment/published_at/confidence. Pencarian teks memakai $regex; belum ada Elasticsearch / text index.")

# 11 Dashboard
code_slide("Bonus 4. Dashboard (Streamlit + Plotly)", "Donut distribusi, tren waktu, kata kunci, volume per layer, comment explorer",
           [snip("dashboard/app.py", 279, 293), snip("dashboard/app.py", 305, 315)], 9,
           "Dashboard di port 8501 (service dashboard di docker-compose.yml, baris 30-48). Data diambil langsung dari koleksi gold_sentiment.")

prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
